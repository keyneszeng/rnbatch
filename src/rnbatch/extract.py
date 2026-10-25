"""正文抽取：docx / doc / pptx / xlsx / pdf / 纯文本 / 图片 OCR。

设计取舍：
- 文档解析用可选依赖，缺哪个降级哪个，不让整个工具挂掉。
- 图片 OCR 走 macOS 原生 Vision（pyobjc 桥），离线、支持中英混排；非 macOS 平台自动跳过。
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

PHOTO_EXT = {".jpg", ".jpeg", ".png", ".heic", ".heif", ".gif", ".webp",
             ".tif", ".tiff", ".bmp", ".raw", ".cr2", ".nef", ".arw", ".dng"}
DOC_EXT = {".doc", ".docx", ".pages", ".rtf", ".txt", ".md"}
SHEET_EXT = {".xls", ".xlsx", ".csv", ".numbers"}
SLIDE_EXT = {".ppt", ".pptx", ".key"}
PDF_EXT = {".pdf"}
VIDEO_EXT = {".mp4", ".mov", ".avi", ".mkv", ".m4v"}
ARCHIVE_EXT = {".zip", ".rar", ".7z"}


def type_of(ext: str) -> str:
    """把扩展名映射成中文类型标签，供 {type} 占位符使用。"""
    ext = ext.lower()
    for label, group in (
        ("照片", PHOTO_EXT), ("视频", VIDEO_EXT), ("PDF", PDF_EXT),
        ("文档", DOC_EXT), ("表格", SHEET_EXT), ("演示", SLIDE_EXT),
        ("压缩包", ARCHIVE_EXT),
    ):
        if ext in group:
            return label
    return "其他"


# ------------------------------------------------------------------ 各格式解析
def _from_docx(p: Path) -> str:
    from docx import Document
    d = Document(str(p))
    parts = [para.text for para in d.paragraphs if para.text.strip()]
    for tbl in d.tables:
        for row in tbl.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _from_pptx(p: Path) -> str:
    from pptx import Presentation
    prs = Presentation(str(p))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"[第{i}页]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    parts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _from_xlsx(p: Path, max_rows: int = 80) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(p), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[工作表 {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                parts.append("…(截断)")
                break
            cells = [str(c) for c in row if c not in (None, "")]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _from_pdf(p: Path, max_pages: int = 4) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(p))
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")           # 有些 PDF 只是空密码加密
        except Exception:
            return "（PDF 已加密，无法读取正文）"
    parts = []
    for i, page in enumerate(reader.pages):
        if i >= max_pages:
            parts.append(f"…(仅取前{max_pages}页)")
            break
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    text = "\n".join(parts).strip()
    if not text:
        return "（PDF 没有文字层，可能是扫描件；需要 OCR 才能读出内容）"
    return text


def _from_plain(p: Path) -> str:
    for enc in ("utf-8", "gb18030", "utf-16", "latin-1"):
        try:
            return p.read_text(encoding=enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return p.read_bytes().decode("utf-8", "replace")


def _from_textutil(p: Path) -> str:
    """老格式（.doc/.rtf/.odt/.html）交给 macOS 自带 textutil。"""
    if sys.platform != "darwin":
        return ""
    try:
        r = subprocess.run(["/usr/bin/textutil", "-convert", "txt", "-stdout", str(p)],
                           capture_output=True, timeout=30)
        if r.returncode == 0:
            return r.stdout.decode("utf-8", "replace")
    except Exception:
        pass
    return ""


def ocr_image(p: Path) -> str:
    """macOS 原生 Vision OCR。离线、不联网、支持 zh-Hans + en-US。"""
    if sys.platform != "darwin":
        return ""
    try:
        import Quartz
        import Vision
        from Foundation import NSURL
    except Exception:
        return ""
    try:
        url = NSURL.fileURLWithPath_(str(p))
        src = Quartz.CGImageSourceCreateWithURL(url, None)
        if src is None:
            return ""
        img = Quartz.CGImageSourceCreateImageAtIndex(src, 0, None)
        if img is None:
            return ""
        collected: list[str] = []

        def handler(request, error):
            for obs in (request.results() or []):
                cands = obs.topCandidates_(1)
                if cands and len(cands):
                    collected.append(cands[0].string())

        req = Vision.VNRecognizeTextRequest.alloc().initWithCompletionHandler_(handler)
        req.setRecognitionLevel_(0)          # 0 = accurate
        req.setUsesLanguageCorrection_(True)
        try:
            req.setRecognitionLanguages_(["zh-Hans", "en-US"])
        except Exception:
            pass
        h = Vision.VNImageRequestHandler.alloc().initWithCGImage_options_(img, None)
        h.performRequests_error_([req], None)
        return "\n".join(collected)
    except Exception as e:
        return f"（OCR 失败：{type(e).__name__}: {e}）"


def _from_image(p: Path) -> str:
    ocr = ocr_image(p).strip()
    meta = []
    try:
        from PIL import Image, ExifTags
        im = Image.open(str(p))
        meta.append(f"尺寸 {im.width}x{im.height}")
        try:
            for tag, val in im.getexif().items():
                name = ExifTags.TAGS.get(tag)
                if name in ("DateTimeOriginal", "DateTime", "Make", "Model"):
                    meta.append(f"{name}: {val}")
        except Exception:
            pass
    except Exception:
        pass
    tail = "，".join(meta)
    if not ocr or ocr.startswith("（OCR"):
        return f"（图片无可识别文字）{tail}"
    return f"[OCR 文字]\n{ocr}\n[图片信息] {tail}"


def extract_text(p: Path, limit: int = 4000) -> str:
    """抽取单个文件的正文，供人或 AI 判断关键词。"""
    p = Path(p)
    ext = p.suffix.lower()
    try:
        if ext == ".docx":
            t = _from_docx(p)
        elif ext == ".pptx":
            t = _from_pptx(p)
        elif ext in (".xlsx", ".xlsm"):
            t = _from_xlsx(p)
        elif ext == ".pdf":
            t = _from_pdf(p)
        elif ext in PHOTO_EXT:
            t = _from_image(p)
        elif ext in (".doc", ".rtf", ".rtfd", ".odt", ".html", ".htm", ".webarchive"):
            t = _from_textutil(p) or _from_plain(p)
        elif ext in (".txt", ".md", ".csv", ".json", ".xml"):
            t = _from_plain(p)
        else:
            t = _from_plain(p) if p.stat().st_size < 3_000_000 else "（二进制文件，未解析）"
    except Exception as e:
        t = f"（解析失败：{type(e).__name__}: {e}）"
    t = re.sub(r"\n{3,}", "\n\n", t).strip()
    return t[:limit]


# 抽正文的优先级：文档最先，照片最后（OCR 最慢，且命中率低于正文）
SOURCE_PRIORITY = {"文档": 0, "PDF": 1, "演示": 2, "表格": 3}

# 同属「文档」时，真正的办公文档比裸文本更可能是正式材料；
# .txt / .md 什么内容都可能是，排在后面。
EXT_PRIORITY = {
    ".docx": 0, ".doc": 1, ".pages": 2, ".rtf": 3,
    ".md": 4, ".txt": 5, ".csv": 6, ".json": 7, ".xml": 8,
}


def preferred_sources(files, n: int = 3):
    """从一堆文件里挑最可能藏着关键词的几个。

    三条排序依据，缺一不可：
    1. 类型 —— 文档 > PDF > 演示 > 表格 > 照片（OCR 慢且命中率低）
    2. 扩展名 —— 正式文档 > 裸文本（见 EXT_PRIORITY）
    3. 体积降序 —— 正文越长，关键词信号越强

    第 3 条是实测补上的：早先只按文件名排序，结果「项目说明.txt」（正文四行）
    被「café.txt」（正文三个字）挤出了候选集，于是候选词里全是
    「报价」「补充材料」这类噪音，真正的项目名反而没被提出来。
    """
    def rank(f):
        try:
            size = f.stat().st_size
        except OSError:
            size = 0
        return (SOURCE_PRIORITY.get(type_of(f.suffix), 5),
                EXT_PRIORITY.get(f.suffix.lower(), 9),
                -min(size, 500_000),
                str(f).lower())
    return sorted(files, key=rank)[:n]


def rich_texts(files, want: int = 3, limit: int = 3000, pool: int = 12):
    """按优先级逐个抽正文，跳过「解析失败 / 没有文字层 / 空文件」，凑够 want 份为止。

    为什么不能只取前 want 个：损坏的 docx、没有文字层的扫描 PDF、只有一行字的
    空壳文件都排在前面，它们占掉名额后整次抽取就颗粒无收。多试几个才有兜底。
    """
    out = []
    for f in preferred_sources(files, max(pool, want)):
        try:
            text = extract_text(f, limit)
        except Exception:
            continue
        if not text.strip():
            continue
        # 抽取层自己的占位提示（（解析失败…）/（PDF 没有文字层…））不算正文
        if text.lstrip().startswith("（"):
            continue
        out.append((f, text))
        if len(out) >= want:
            break
    return out
